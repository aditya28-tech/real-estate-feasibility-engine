"""
engine/export.py  -  STEP 9: one-click EXCEL REPORT of everything the app calculates.

build_excel_report(project, res, stress=None, mc=None) -> bytes of an .xlsx file with sheets:
    Summary · Inputs · P&L & Metrics · Cash Flow · Collections · Timeline · Stress Test · Monte Carlo
The numbers are values (not formulas) so anyone can open, read, print or forward the file without the app.
A hidden sheet keeps the full project, so the report can be opened again in the app (project_from_excel).
"""
from __future__ import annotations

import io
import json
from datetime import datetime

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as GL

from .inputs import CRORE, Project, from_dict, to_dict
from .report import key_metrics_table, kpi_cards, takeaways

NAVY, F = "1F3864", "Arial"
HEAD = PatternFill("solid", fgColor="D9E1F2"); SEC = PatternFill("solid", fgColor=NAVY); TOT = PatternFill("solid", fgColor="F2F2F2")
GOOD = PatternFill("solid", fgColor="E2F0D9"); REVIEW = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CR, PCT, INT, QF = '#,##0.00;(#,##0.00);"-"', '0.0%;(0.0%);"-"', '#,##0;(#,##0);"-"', '"Q"0'
JSON_SHEET = "_project"


class _Sheet:
    """Tiny helper that writes rows top-down with consistent styling."""
    def __init__(self, wb, title, widths=(34,) + (14,) * 45):
        self.ws = wb.create_sheet(title)
        self.ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths): self.ws.column_dimensions[GL(i + 1)].width = w
        self.r = 1

    def title(self, text, sub=""):
        c = self.ws.cell(self.r, 1, text); c.font = Font(name=F, size=16, bold=True, color=NAVY); self.r += 1
        if sub: self.ws.cell(self.r, 1, sub).font = Font(name=F, size=9, italic=True, color="7F7F7F"); self.r += 1
        self.r += 1

    def section(self, text, width=10):
        for col in range(1, width + 1): self.ws.cell(self.r, col).fill = SEC
        self.ws.cell(self.r, 1, text).font = Font(name=F, size=11, bold=True, color="FFFFFF"); self.r += 1

    def header(self, cols, start=1):
        for i, h in enumerate(cols):
            c = self.ws.cell(self.r, start + i, h); c.font = Font(name=F, bold=True, size=10); c.fill = HEAD; c.border = BOX
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        self.r += 1

    def row(self, values, fmts=None, bold=False, fill=None, start=1):
        for i, v in enumerate(values):
            if isinstance(v, (np.floating, np.integer)): v = v.item()
            if isinstance(v, float) and np.isnan(v): v = None
            c = self.ws.cell(self.r, start + i, v); c.font = Font(name=F, size=10, bold=bold); c.border = BOX
            if fmts and i < len(fmts) and fmts[i]: c.number_format = fmts[i]
            if fill: c.fill = fill
        self.r += 1
        return self.r - 1

    def gap(self, n=1): self.r += n

    def note(self, text):
        self.ws.cell(self.r, 1, text).font = Font(name=F, size=9, italic=True, color="7F7F7F"); self.r += 1


def build_excel_report(project: Project, res: dict, stress: dict | None = None, mc: dict | None = None) -> bytes:
    tl, cv, co, cr, cf = res["timeline"], res["curves"], res["collections"], res["costs"], res["cashflow"]
    pnl, m, info = cf.pnl, cf.metrics, project.info
    stamp = datetime.now().strftime("%d %b %Y %H:%M")
    wb = Workbook(); wb.remove(wb.active)
    n_q = max(tl.last_q, 1)

    # ------------------------------------------------------------------ 1. Summary
    s = _Sheet(wb, "Summary", (30, 16, 46, 12, 30, 16, 46, 12))
    s.title(f"{info.name} — Feasibility report",
            f"{info.location} · {info.project_type} · {info.deal_type} deal · scenario: {info.scenario} · generated {stamp} · values, not formulas")
    status = "✔ ALL CHECKS PASSED" if not res["problems"] else f"✖ {len(res['problems'])} CHECK(S) FAILED – review before use"
    c = s.ws.cell(s.r, 1, status); c.font = Font(name=F, bold=True, size=12, color="006100" if not res["problems"] else "C00000"); s.r += 2
    s.section("KEY PERFORMANCE INDICATORS (green = beats benchmark, amber = review)", 8)
    s.header(["KPI", "Value", "Detail", "", "KPI", "Value", "Detail", ""])
    cards = kpi_cards(project, res)
    for i in range(0, 12, 2):
        a, b = cards[i], cards[i + 1]
        rr = s.row([a["label"], a["value"], a["sub"], "", b["label"], b["value"], b["sub"], ""])
        for col, card in ((2, a), (6, b)):
            s.ws.cell(rr, col).fill = {"good": GOOD, "review": REVIEW}.get(card["status"], PatternFill())
            s.ws.cell(rr, col).font = Font(name=F, bold=True, size=11)
    s.gap()
    s.section("KEY TAKEAWAYS", 8)
    for i, t in enumerate(takeaways(project, res), 1):
        s.ws.cell(s.r, 1, f"{i}. {t}").alignment = Alignment(wrap_text=True, vertical="top")
        s.ws.merge_cells(start_row=s.r, start_column=1, end_row=s.r, end_column=8); s.ws.row_dimensions[s.r].height = 30; s.r += 1
    s.gap()
    s.section("PROFIT & LOSS (₹ Crore)", 8)
    s.header(["", "₹ Crore", "₹ / sq ft", "% of revenue"])
    sa = cr.total_sa or 1
    for lab, v, b in [("Total revenue (GDV)", pnl["revenue"], True), ("Total cost", pnl["total_cost"], False), ("EBITDA", pnl["ebitda"], True),
                      (f"Tax @ {pnl['tax_rate']:.1%} ({info.entity})", pnl["tax"], False), ("PAT", pnl["pat"], True)]:
        s.row([lab, v, v * CRORE / sa, v / pnl["revenue"] if pnl["revenue"] else 0], [None, CR, INT, PCT], bold=b)
    if problems := res["problems"]:
        s.gap(); s.section("CHECKS THAT FAILED", 8)
        for p_ in problems: s.note(p_)
    # cost breakdown chart
    s.gap(); top = s.r
    groups = [("Land / landowner", cr.amounts["land"] + cr.amounts["sd"] + cr.amounts["brok"]), ("Construction", cr.construction_heads_total),
              ("Marketing", cr.amounts["mkt"]), ("Overheads", cr.amounts["ovh"]), ("Legal & BAR", cr.amounts["legal"] + cr.amounts["bar"])]
    groups.append(("Other", cr.total_cost - sum(v for _, v in groups)))
    s.header(["Cost group", "₹ Crore"])
    g0 = s.r
    for lab, v in groups: s.row([lab, v], [None, CR])
    ch = BarChart(); ch.type = "bar"; ch.title = "Cost breakdown (₹ Cr)"; ch.legend = None; ch.height = 7; ch.width = 16
    ch.add_data(Reference(s.ws, min_col=2, min_row=g0, max_row=s.r - 1)); ch.set_categories(Reference(s.ws, min_col=1, min_row=g0, max_row=s.r - 1))
    ch.x_axis.delete = False; ch.y_axis.delete = False
    s.ws.add_chart(ch, f"E{top}")
    s.ws.page_setup.orientation = "landscape"; s.ws.page_setup.fitToWidth = 1; s.ws.page_setup.fitToHeight = 0
    s.ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ------------------------------------------------------------------ 2. Inputs
    s = _Sheet(wb, "Inputs", (34, 22, 16, 16, 14, 14, 14, 12, 12, 14, 24, 12, 12, 14, 12, 44))
    s.title("Inputs used for this report", "Exactly as entered in the app (percentages shown as %)")
    s.section("PROJECT & LAND", 6)
    L = project.land
    for lab, v, f in [("Project name", info.name, None), ("Location", info.location, None), ("Project type", info.project_type, None),
                      ("Project start (Q1)", info.start_date, "dd-mmm-yyyy"), ("Deal type", info.deal_type, None), ("Taxable entity", info.entity, None),
                      ("Active scenario", info.scenario, None), ("Land area (sq yd)", L.area_sqyd, INT), ("Market rate (₹/sq yd)", L.market_rate, INT),
                      ("DLC rate (₹/sq m)", L.dlc_rate_sqm, INT), ("Landowner share (JV)", L.lo_share, PCT), ("Stamp duty JV (% of DLC)", L.stamp_duty_jv, '0.00%'),
                      ("Purchase rate – outright (₹/sq yd)", L.purchase_rate, INT), ("Stamp duty – outright", L.stamp_duty_outright, '0.00%'),
                      ("Land payment curve / start / end", f"{L.payment_curve} Q{L.payment_start_q}–Q{L.payment_end_q}", None),
                      ("Catch-up lag (quarters)", project.catch_up_lag, '0')]:
        s.row([lab, v], [None, f])
    s.gap(); s.section("SCENARIOS & ZONES", 6)
    s.header(["Scenario", "Sales period ×", "Price change", "Cost change", "Delay (Q)"])
    for sc in project.scenarios: s.row([sc.name, sc.sales_period_mult, sc.price_change, sc.cost_change, sc.delay_q], [None, '0.00"x"', PCT, PCT, '0'])
    s.header(["Zone", "Share of land"])
    for z in project.zones: s.row([z.name, z.share], [None, PCT])
    s.gap(); s.section("PRODUCTS (with final timing)", 16)
    s.header(["Product", "Zone", "Category", "Included", "Area (sq ft)", "Rate ₹/sq ft", "Constr. ₹/sq ft", "Sales period", "Sales curve",
              "Constr. qtrs", "Payment plan", "Launch", "Sales end", "Constr. start", "Constr. end", "How timing was decided"])
    for p in project.products:
        t = tl.products.get(p.name)
        s.row([p.name, p.zone, p.category, "Yes" if p.include else "No", p.area_sqft, p.rate, p.construction_rate, p.sales_period_q, p.sales_curve,
               p.construction_duration_q, p.clp_plan, t.launch if t else None, t.sales_end if t else None, t.cons_start if t else None,
               t.cons_end if t else None, t.note if t else "not active"], [None, None, None, None, INT, INT, INT, '0', None, '0', None, QF, QF, QF, QF, None])
    s.gap(); s.section("COST HEADS", 10)
    s.header(["Cost head", "Input 1", "Input 2 / basis", "Applicable", "Phasing curve", "Start Q", "End Q", "Amount ₹ Cr", "₹ / sq ft", "% of GDV"])
    for c_ in project.costs:
        a = cr.amounts[c_.key]
        s.row([c_.name, c_.rate1, c_.rate2 if c_.rate2 is not None else c_.basis, "Yes" if c_.applicable else "No", c_.curve, c_.start_q, c_.end_q,
               a, a * CRORE / sa, a / cr.total_revenue if cr.total_revenue else 0], [None, '#,##0.####', None, None, None, '0', '0', CR, INT, '0.00%'])
    s.row(["TOTAL COST", None, None, None, None, None, None, cr.total_cost, cr.total_cost * CRORE / sa, cr.total_cost / cr.total_revenue],
          [None] * 7 + [CR, INT, '0.00%'], bold=True, fill=TOT)
    s.gap(); s.section("CALCULATORS", 6)
    b = cr.bar
    s.row(["BAR / premium FAR", f"{project.bar.authority} / {project.bar.land_use}: extra {b['extra_sqm']:,.0f} sq m × ₹{project.bar.reserved_price_sqm:,.0f} × {b['pct']:.0%}", None, None, None, None, None, b["cost"]],
          [None] * 7 + [CR])
    s.row(["AIF interest (if switched on)", f"{project.aif.rate:.0%} for {project.aif.tenure_years} yrs, {project.aif.compounding}", None, None, None, None, None, cr.aif["cost"]], [None] * 7 + [CR])
    s.header(["Initial investment item (JV)", "₹ Crore", "Payout quarter"])
    for k_, v in cr.initial_items.items(): s.row([k_, v, project.initial.payout_q.get(k_)], [None, CR, QF])
    s.row(["Initial investment used (" + info.deal_type + ")", cr.initial_used], [None, CR], bold=True, fill=TOT)
    s.gap(); s.section("PAYMENT PLANS (CLP)", 6)
    s.header(["Plan / milestone", "Trigger", "Value", "% of sale", "Used by"])
    for pl in project.plans:
        s.row([pl.name, "", "", sum(m_.pct for m_ in pl.milestones), ", ".join(p.name for p in project.products if p.clp_plan == pl.name)],
              [None, None, None, PCT], bold=True, fill=TOT)
        for m_ in pl.milestones:
            s.row(["   " + m_.name, m_.trigger, m_.value, m_.pct], [None, None, '0' if m_.trigger.startswith("Booking") else '0%', PCT])
    s.gap(); s.section("TAX, COST OF EQUITY & BENCHMARKS", 6)
    fz = project.finance
    for lab, v, f in [("Tax – Company", fz.tax_company, PCT), ("Tax – LLP", fz.tax_llp, PCT), ("Cost of equity", fz.cost_of_equity, PCT)] + \
                     [(f"Benchmark – {k_.replace('_', ' ')}", v, '0.00"x"' if k_ == "equity_multiple" else PCT) for k_, v in fz.benchmarks.items()]:
        s.row([lab, v], [None, f])
    s.ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ 3. P&L & Metrics
    s = _Sheet(wb, "P&L & Metrics", (40, 16, 16, 16, 16))
    s.title("Revenue, cost, P&L and key metrics")
    s.section("A. REVENUE STATEMENT", 5); s.header(["Product", "₹ Crore", "₹ / sq ft", "% of total", "Category"])
    sc_ = project.scenario()
    for p in project.active_products():
        s.row([p.name, cv.revenue[p.name], p.rate * (1 + sc_.price_change), cv.revenue[p.name] / cr.total_revenue, p.category], [None, CR, INT, PCT])
    s.row(["TOTAL REVENUE", cr.total_revenue, cr.total_revenue * CRORE / sa, 1.0], [None, CR, INT, PCT], bold=True, fill=TOT)
    s.gap(); s.section("B. COST STATEMENT", 5); s.header(["Cost head", "₹ Crore", "₹ / sq ft", "% of revenue", "Applicable"])
    for c_ in project.costs:
        a = cr.amounts[c_.key]
        s.row([c_.name, a, a * CRORE / sa, a / cr.total_revenue if cr.total_revenue else 0, "Yes" if c_.applicable else "No"], [None, CR, INT, '0.00%'])
    s.row(["TOTAL COST", cr.total_cost, cr.total_cost * CRORE / sa, cr.total_cost / cr.total_revenue], [None, CR, INT, PCT], bold=True, fill=TOT)
    s.gap(); s.section("C. PROFIT & LOSS", 5); s.header(["", "₹ Crore", "₹ / sq ft", "% of revenue"])
    for lab, v in [("Revenue", pnl["revenue"]), ("Total cost", pnl["total_cost"]), ("EBITDA", pnl["ebitda"]), ("Tax", pnl["tax"]), ("PAT", pnl["pat"]),
                   ("Developer's revenue share", pnl["developer_revenue"]), ("Developer's cost (excl. landowner share)", pnl["developer_cost"])]:
        s.row([lab, v, v * CRORE / sa, v / pnl["revenue"] if pnl["revenue"] else 0], [None, CR, INT, PCT], bold=lab in ("EBITDA", "PAT"))
    s.row(["EBITDA margin on developer's share", pnl["ebitda_margin_on_dev_share"]], [None, PCT])
    s.gap(); s.section("D. JV REVENUE SPLIT", 5); s.header(["Party", "Share", "₹ Crore"])
    s.row(["Developer", cr.dev_share, pnl["developer_revenue"]], [None, PCT, CR]); s.row(["Landowner", 1 - cr.dev_share, pnl["landowner_revenue"]], [None, PCT, CR])
    s.gap(); s.section("E. KEY INVESTMENT METRICS", 5); s.header(["Metric", "Value", "Benchmark", "Status"])
    for row_ in key_metrics_table(project, res):
        rr = s.row([row_["metric"], row_["value"], row_["benchmark"], row_["status"]])
        if row_["status"]: s.ws.cell(rr, 4).fill = GOOD if row_["status"] == "Strong" else REVIEW

    # ------------------------------------------------------------------ 4. Cash Flow
    s = _Sheet(wb, "Cash Flow", (44,) + (12,) * (n_q + 2))
    s.title("Quarterly cash flow (₹ Crore)", "Columns = quarters from the project start. Negative = cash out.")
    hdr = ["Quarter"] + [f"Q{q}" for q in range(1, n_q + 1)] + ["Total"]
    s.header(hdr)
    s.row(["Quarter end"] + tl.end_dates[:n_q], [None] + ["mmm-yy"] * n_q)
    s.row(["Financial year"] + tl.fy[:n_q])
    s.row(["Phase"] + tl.phase[:n_q])
    def qrow(label, arr, bold=False, fill=None, sign=1.0):
        vals = [sign * float(v) for v in arr[:n_q]]
        return s.row([label] + vals + [sum(vals)], [None] + [CR] * (n_q + 1), bold=bold, fill=fill)
    s.section("INFLOWS", n_q + 2)
    qrow("Customer collections", cf.inflow, True)
    qrow("   of which landowner share earned", co.landowner); qrow("   of which developer share", co.developer)
    s.section("OUTFLOWS", n_q + 2)
    for c_ in project.costs:
        if abs(cf.outflows[c_.key]).sum() > 1e-9: qrow("   " + c_.name, cf.outflows[c_.key], sign=-1)
    qrow("Total costs paid", cf.total_outflow, True, TOT, -1)
    qrow("Advance to landowner", cf.lo_advance, sign=-1); qrow("Cash paid to landowner (incl. advance)", cf.lo_paid, sign=-1)
    s.section("NET CASH FLOW", n_q + 2)
    qrow("Project net cash flow (before landowner)", cf.project_cf, True)
    r_dcf = qrow("DEVELOPER NET CASH FLOW (pre-tax)", cf.developer_cf, True, TOT)
    r_cum = s.row(["Developer cumulative (funding curve)"] + list(np.cumsum(cf.developer_cf)[:n_q]), [None] + [CR] * n_q, bold=True)
    qrow("Tax (indicative)", cf.tax, sign=-1); qrow("Developer net cash flow (post-tax, indicative)", cf.developer_cf_post_tax)
    s.gap(); s.section("YEARLY SUMMARY BY FINANCIAL YEAR", 6)
    s.header(["Financial year", "Collections", "Costs paid", "Paid to landowner", "Developer net", "Developer cumulative"])
    run = 0.0
    for fy, v in cf.fy_summary.items():
        if abs(v["collections"]) + abs(v["costs_paid"]) + abs(v["paid_to_landowner"]) < 1e-9: continue
        run += v["developer_net"]
        s.row([fy, v["collections"], v["costs_paid"], v["paid_to_landowner"], v["developer_net"], run], [None] + [CR] * 5)
    bar = BarChart(); bar.title = "Developer cash flow (₹ Cr) and funding curve"; bar.height = 8; bar.width = 28
    bar.add_data(Reference(s.ws, min_col=2, max_col=n_q + 1, min_row=r_dcf), from_rows=True, titles_from_data=False)
    bar.set_categories(Reference(s.ws, min_col=2, max_col=n_q + 1, min_row=4))
    ln = LineChart(); ln.add_data(Reference(s.ws, min_col=2, max_col=n_q + 1, min_row=r_cum), from_rows=True, titles_from_data=False)
    ln.y_axis.axId = 200; ln.y_axis.crosses = "max"; bar += ln
    for ax in (bar.x_axis, bar.y_axis, ln.y_axis): ax.delete = False
    bar.legend = None
    s.ws.add_chart(bar, f"A{s.r + 2}")
    s.ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ 5. Collections
    s = _Sheet(wb, "Collections", (44,) + (12,) * (n_q + 2))
    s.title("Bookings, collections and payment-plan milestones (₹ Crore)")
    s.header(["Quarter"] + [f"Q{q}" for q in range(1, n_q + 1)] + ["Total"])
    for name in co.collections:
        s.section(name, n_q + 2)
        qrow("   Bookings (sales value)", co.bookings[name]); qrow("   Booking-linked cash", co.booking_linked[name])
        qrow("   Construction-linked cash", co.construction_linked[name]); qrow("   Collections", co.collections[name], True)
        s.row(["   Receivables (cumulative)"] + list(co.receivables[name][:n_q]), [None] + [CR] * n_q)
    s.section("PROJECT", n_q + 2)
    qrow("Total collections", co.total, True, TOT); qrow("Landowner share", co.landowner); qrow("Developer share", co.developer)
    s.gap(); s.section("MILESTONE SCHEDULE", 5); s.header(["Product", "Payment plan", "Milestone", "% of sale", "Falls due"])
    for x in co.schedule: s.row([x.product, x.plan, x.name, x.pct, x.text], [None, None, None, PCT])
    s.ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ 6. Timeline
    s = _Sheet(wb, "Timeline", (34,) + (5,) * n_q)
    s.title("Project timeline", "S = selling · C = construction · # = both")
    s.header(["Product"] + [f"Q{q}" for q in range(1, n_q + 1)])
    for t in tl.products.values():
        vals = ["#" if (t.launch <= q <= t.sales_end and t.cons_start <= q <= t.cons_end) else "S" if t.launch <= q <= t.sales_end
                else "C" if t.cons_start <= q <= t.cons_end else "" for q in range(1, n_q + 1)]
        rr = s.row([t.name] + vals)
        for j, v in enumerate(vals):
            if v: s.ws.cell(rr, 2 + j).fill = PatternFill("solid", fgColor={"S": "A9D08E", "C": "9BC2E6", "#": "8EA9DB"}[v])
    s.row(["Phase"] + tl.phase[:n_q])
    s.gap()
    for t in tl.products.values(): s.note(f"{t.name}: {t.note}")

    # ------------------------------------------------------------------ 7. Stress Test
    s = _Sheet(wb, "Stress Test", (34, 16, 16, 16, 16, 16, 16, 16))
    s.title("Stress test", "Each line is a full model run with the lever shown (on top of the active scenario).")
    if stress:
        s.section("SCENARIO COMPARISON", 8)
        s.header(["Scenario", "GDV ₹Cr", "EBITDA ₹Cr", "PAT ₹Cr", "Developer IRR", "Dev NPV ₹Cr", "Peak ₹Cr", "Payback"])
        for x in stress["scenarios"]:
            s.row([x["scenario"], x["gdv"], x["ebitda"], x["pat"], x["developer_irr"], x["developer_npv"], x["peak"], f"Q{x['payback_q']}" if x["payback_q"] else "never"],
                  [None, CR, CR, CR, PCT, CR, CR])
        names = {"price": ("Sales price shock", '+0%;-0%;0%'), "cost": ("Construction cost shock", '+0%;-0%;0%'),
                 "speed": ("Sales period multiplier", '0.00"x"'), "delay": ("Extra delay (quarters)", '"+"0" Q"')}
        for lv, rows in stress["one_way"].items():
            s.gap(); s.section(f"ONE-WAY: {names[lv][0].upper()}", 8)
            s.header([names[lv][0], "Developer IRR", "Dev NPV ₹Cr", "PAT ₹Cr", "EBITDA ₹Cr", "Peak ₹Cr", "Payback"])
            for x in rows:
                s.row([x["value"], x["developer_irr"], x["developer_npv"], x["pat"], x["ebitda"], x["peak"], f"Q{x['payback_q']}" if x["payback_q"] else "never"],
                      [names[lv][1], PCT, CR, CR, CR, CR])
        for key, title, rows_, cols_, rf, cf_, vf in [
                ("irr_grid", "DEVELOPER IRR – construction cost (down) × price (across)", [0, .05, .1, .15, .2], [-.1, -.05, 0, .05, .1], '+0%;-0%;0%', '+0%;-0%;0%', PCT),
                ("pat_grid", "PAT ₹Cr – construction cost (down) × price (across)", [0, .05, .1, .15, .2], [-.1, -.05, 0, .05, .1], '+0%;-0%;0%', '+0%;-0%;0%', CR),
                ("peak_grid", "PEAK INVESTMENT ₹Cr – sales period (down) × delay (across)", [.75, 1, 1.25, 1.5, 2], [0, 1, 2, 4], '0.00"x"', '"+"0" Q"', CR)]:
            s.gap(); s.section(title, 8)
            s.row([""] + cols_, [None] + [cf_] * len(cols_), bold=True, fill=HEAD)
            for i, rv in enumerate(rows_): s.row([rv] + list(stress[key][i]), [rf] + [vf] * len(cols_))
        h = stress["headroom"]
        s.gap(); s.section("HEADROOM", 8)
        for lab, v in [("Price shock at which EBITDA = 0", h["price_to_zero_ebitda"]), ("Cost overrun at which EBITDA = 0", h["cost_to_zero_ebitda"]),
                       (f"Price shock at which developer IRR = {h['hurdle']:.0%}", h["price_to_hurdle"]),
                       (f"Cost overrun at which developer IRR = {h['hurdle']:.0%}", h["cost_to_hurdle"])]:
            s.row([lab, v if v is not None else "beyond tested range"], [None, '+0.0%;-0.0%'])
        s.gap(); s.section("WHAT MOVES DEVELOPER IRR THE MOST", 8); s.header(["Lever", "Downside case", "IRR", "Upside case", "IRR", "Change (downside)"])
        for x in stress["tornado"]:
            s.row([x["lever"], x["down_value"], x["down"], x["up_value"], x["up"], x["impact"]], [None, '0.00', PCT, '0.00', PCT, '+0.0%;-0.0%'])
    else:
        s.note("Stress test not included in this report.")

    # ------------------------------------------------------------------ 8. Monte Carlo
    s = _Sheet(wb, "Monte Carlo", (34, 16, 16, 16, 16, 16, 16, 16))
    s.title("Monte Carlo simulation")
    if mc:
        s.section("RESULTS", 6)
        for lab, v, f in [("Number of runs", mc["n"], INT), ("Random seed", mc["seed"], '0'), ("Runs that broke a model check", mc["failed_runs"], '0'),
                          (f"P(developer IRR ≥ {mc['hurdle']:.0%})", mc["p_irr_above_hurdle"], PCT), ("P(developer NPV > 0)", mc["p_npv_positive"], PCT),
                          ("P(loss: PAT < 0)", mc["p_loss"], PCT)]:
            s.row([lab, v], [None, f], bold=lab.startswith("P("))
        s.header(["", "P10 (pessimistic)", "P50 (median)", "P90 (optimistic)"])
        s.row(["Developer IRR", mc["irr_p10"], mc["irr_p50"], mc["irr_p90"]], [None, PCT, PCT, PCT])
        s.row(["Developer NPV ₹Cr", mc["npv_p10"], mc["npv_p50"], mc["npv_p90"]], [None, CR, CR, CR])
        s.row(["PAT ₹Cr", mc["pat_p10"], mc["pat_p50"], mc["pat_p90"]], [None, CR, CR, CR])
        s.row(["Peak investment ₹Cr (P50 / P90 / worst)", mc["peak_p50"], mc["peak_p90"], mc["peak_max"]], [None, CR, CR, CR])
        s.gap(); s.section("PROBABILITY THAT DEVELOPER IRR IS AT LEAST …", 6)
        irr_ok = mc["irr"][~np.isnan(mc["irr"])]
        s.header(["IRR at least", "Probability"])
        p0 = s.r
        for thr in [0.0, 0.10, 0.18, 0.25, 0.40, 0.60, 0.80, 1.00, 1.25]:
            s.row([thr, float(np.mean(irr_ok >= thr)) if len(irr_ok) else 0.0], [PCT, PCT])
        ch = LineChart(); ch.title = "Probability IRR ≥ X"; ch.height = 7; ch.width = 14; ch.legend = None
        ch.add_data(Reference(s.ws, min_col=2, min_row=p0, max_row=s.r - 1)); ch.set_categories(Reference(s.ws, min_col=1, min_row=p0, max_row=s.r - 1))
        ch.x_axis.delete = False; ch.y_axis.delete = False
        s.ws.add_chart(ch, f"D{p0 - 1}")
        s.gap(4); s.section("EVERY RUN (inputs drawn and results)", 8)
        s.header(["Run", "Price shock", "Cost shock", "Sales period ×", "Delay (Q)", "Developer IRR", "Dev NPV ₹Cr", "Peak ₹Cr"])
        sm = mc["samples"]
        for i in range(mc["n"]):
            s.row([i + 1, sm["price"][i], sm["cost"][i], sm["speed"][i], int(sm["delay"][i]), mc["irr"][i], mc["npv"][i], mc["peak"][i]],
                  ['0', PCT, PCT, '0.00', '0', PCT, CR, CR])
    else:
        s.note("Monte Carlo not run before this report was created (run it on the 'S · Stress & Monte Carlo' screen, then download again).")

    # ------------------------------------------------------------------ hidden project copy (to reopen in the app)
    ws = wb.create_sheet(JSON_SHEET)
    text = json.dumps(to_dict(project), ensure_ascii=False)
    for i in range(0, len(text), 30000):
        ws.cell(i // 30000 + 1, 1, text[i:i + 30000])
    ws.sheet_state = "hidden"

    wb.active = 0
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def project_from_excel(data: bytes) -> Project:
    """Re-open a report produced by build_excel_report in the app."""
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb[JSON_SHEET]
    text = "".join(r[0] for r in ws.iter_rows(min_col=1, max_col=1, values_only=True) if r[0])
    return from_dict(json.loads(text))
