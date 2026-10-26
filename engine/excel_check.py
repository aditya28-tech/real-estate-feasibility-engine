"""
engine/excel_check.py - reads the Excel FINAL model and compares it with the Python engine.
Works by LABELS (not cell addresses), so it is robust to row positions.
"""
from __future__ import annotations
import numpy as np
import openpyxl

NQ = 40
QCOL0 = 4      # quarter columns start at column D in TIMELINE / CURVES


def _row_values(ws, r):
    return np.array([ws.cell(r, c).value or 0 for c in range(QCOL0, QCOL0 + NQ)], dtype=float)


def read_excel_curves(path: str) -> dict:
    """Returns {'sales': {product: arr}, 'cons': {...}, 'cost': {cost name: arr}, 'profiles': {...}}"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    ws = wb["CURVES"]
    out = {"sales": {}, "cons": {}, "cost": {}, "cost_order": [], "profiles": {}}
    block, section = None, None
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if not isinstance(label, str):
            continue
        if label.startswith("B.  PRODUCT"):   section = "products"; continue
        if label.startswith("C.  AGGREGATE"): section = "profiles"; continue
        if label.startswith("D.  COST HEAD"): section = "costs"; continue
        if label.startswith("E.  CURVE CHECKS"): break
        if section == "products":
            if label.startswith("   Sales % (curve"):         out["sales"][block] = _row_values(ws, r)
            elif label.startswith("   Construction % (curve"): out["cons"][block] = _row_values(ws, r)
            elif not label.startswith("   "):                   block = label
        elif section == "profiles":
            out["profiles"][label] = _row_values(ws, r)
        elif section == "costs":
            if label.startswith("   Spend %"): out["cost"][block] = _row_values(ws, r); out["cost_order"].append(block)
            elif not label.startswith("   "):   block = label
    return out


def read_excel_timeline(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["INPUTS"]
    prods = {}
    for r in range(1, ws.max_row + 1):          # product table: active rows have quarters in P..S
        if ws.cell(r, 5).value == "Yes" and isinstance(ws.cell(r, 16).value, (int, float)) and isinstance(ws.cell(r, 2).value, str):
            prods[ws.cell(r, 2).value] = tuple(int(ws.cell(r, c).value) for c in (16, 17, 18, 19))
    return prods


def read_excel_collections(path: str) -> dict:
    """Bookings, collections, totals, JV split and milestone due quarters from the COLLECTIONS sheet."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["COLLECTIONS"]
    out = {"bookings": {}, "collections": {}, "triggers": {}, "total": None, "landowner": None, "developer": None}
    section, block = None, None
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if not isinstance(label, str):
            continue
        if label.startswith("A.  MILESTONE"): section = "A"; continue
        if label.startswith("B.  PRODUCT"):   section = "B"; continue
        if label.startswith("C.  PROJECT"):   section = "C"; continue
        if label.startswith("D.  COLLECTION"): break
        if section == "A":
            if "   →   " in label:
                block = label.split("   →   ")[0]; out["triggers"][block] = []
            elif block and label.strip() and not label.startswith("   ") and ws.cell(r, 5).value:
                due = ws.cell(r, 6).value
                out["triggers"][block].append((label, int(due) if isinstance(due, (int, float)) else None))
        elif section == "B":
            if label.startswith("   Bookings (sales"):        out["bookings"][block] = _row_values(ws, r)
            elif label.startswith("   Collections (cash"):    out["collections"][block] = _row_values(ws, r)
            elif not label.startswith("   "):                 block = label
        elif section == "C":
            if label.startswith("Total collections (cash"):   out["total"] = _row_values(ws, r)
            elif label.startswith("   Landowner share"):      out["landowner"] = _row_values(ws, r)
            elif label.startswith("   Developer share"):      out["developer"] = _row_values(ws, r)
    return out


def read_excel_cashflow(path: str) -> dict:
    """Quarterly cash rows from the CASHFLOW sheet and P&L / key metrics from the OUTPUT sheet."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["CASHFLOW"]
    out = {"outflows": [], "rows": {}, "pnl": {}, "costs": [], "metrics": {}}
    want = {"Customer collections (100%)": "inflow", "Total outflows (excl. landowner share)": "total_outflow",
            "Cash paid to landowner this quarter": "lo_paid", "Developer net cash flow (pre-tax)": "developer_cf",
            "Developer net cash flow (post-tax, indicative)": "developer_cf_post_tax"}
    section = None
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if not isinstance(label, str):
            continue
        if label.startswith("B.  CASH OUTFLOWS"): section = "B"; continue
        if label.startswith("C.  LANDOWNER"): section = None
        if label.startswith("Project net cash flow"): out["rows"]["project_cf"] = _row_values(ws, r)
        if section == "B" and not label.startswith("Total outflows"):
            out["outflows"].append(_row_values(ws, r))
        if label in want: out["rows"][want[label]] = _row_values(ws, r)
        if label.startswith("Total outflows"): section = None
    ws = wb["OUTPUT"]
    section = None
    for r in range(1, ws.max_row + 1):
        title = ws.cell(r, 2).value
        if isinstance(title, str):
            if title.startswith("B.  COST STATEMENT"): section = "B"
            elif title.startswith("C.  PROFIT"): section = "C"
            elif title.startswith("D.  JV"): section = "D"
            elif title.startswith("E.  KEY"): section = "E"
            continue
        label, val = ws.cell(r, 3).value, ws.cell(r, 4).value
        if not isinstance(label, str):
            continue
        if section == "B" and label not in ("Cost head",) and not label.startswith("TOTAL") and not label.startswith("Total development"):
            out["costs"].append(val or 0.0)
        if label == "SUM TOTAL REVENUE (A)": out["pnl"]["revenue"] = val
        if label == "TOTAL COST (B)": out["pnl"]["total_cost"] = val
        if label == "EBITDA (A − B)": out["pnl"]["ebitda"] = val
        if label == "Tax": out["pnl"]["tax"] = val
        if label == "PROFIT AFTER TAX (PAT)": out["pnl"]["pat"] = val
        if section == "E" and label[:3].strip().rstrip(".").isdigit():
            out["metrics"][label] = val
    return out


def read_excel_stress(path: str) -> dict:
    """Scenario comparison, one-way tables (C1-C4), heat maps (D1-D3) and headroom from the STRESS sheet."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["STRESS"]
    out, section, mode, rows = {}, None, None, []
    val = lambda r, c: ws.cell(r, c).value
    r = 1
    while r <= ws.max_row:
        b = val(r, 2)
        if isinstance(b, str) and b[:3] in ("B. ", "C1.", "C2.", "C3.", "C4.", "D1.", "D2.", "D3.", "E. "):
            section = b.split(".")[0]
            if section in ("D1", "D2", "D3"):
                hdr = r + 2
                cols = [val(hdr, c) for c in range(3, 13) if isinstance(val(hdr, c), (int, float))]
                rr, grid, rowv = hdr + 1, [], []
                while isinstance(val(rr, 2), (int, float)):
                    rowv.append(val(rr, 2)); grid.append([val(rr, 3 + j) for j in range(len(cols))]); rr += 1
                out[section] = dict(row_values=rowv, col_values=cols, grid=np.array(grid, dtype=float))
                r = rr; continue
            if section == "E":
                out["headroom"] = dict(price_to_zero_ebitda=val(r + 1, 3), cost_to_zero_ebitda=val(r + 2, 3))
                break
            hdr = r + 2                       # 'Current model' row, then header row, then selector row
            names = [val(hdr, c) for c in range(3, 13) if val(hdr, c)]
            rr, data = hdr + 2, []
            while isinstance(val(rr, 2), (int, float)):
                data.append(dict(value=val(rr, 2), **{n: val(rr, 3 + j) for j, n in enumerate(names)})); rr += 1
            out[section] = data
            r = rr; continue
        r += 1
    return out
