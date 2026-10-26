"""STEP 9 – Excel report: builds, contains the right numbers, and can be reopened in the app."""
import io, os, sys

import numpy as np
import pytest
from openpyxl import load_workbook

from engine.export import build_excel_report, project_from_excel
from engine.inputs import to_dict
from engine.stress import monte_carlo
from conftest import variant
from engine.cashflow import run_model

SHEETS = ["Summary", "Inputs", "P&L & Metrics", "Cash Flow", "Collections", "Timeline", "Stress Test", "Monte Carlo", "_project"]


def _find(ws, label, col=1):
    for row in ws.iter_rows():
        if row[col - 1].value == label:
            return row
    raise AssertionError(f"'{label}' not found in {ws.title}")


def test_report_builds_with_all_sheets(project, base):
    wb = load_workbook(io.BytesIO(build_excel_report(project, base)))
    assert wb.sheetnames == SHEETS and wb["_project"].sheet_state == "hidden"


def test_report_numbers_match_the_engine(project, base):
    wb = load_workbook(io.BytesIO(build_excel_report(project, base)))
    pnl = _find(wb["Summary"], "PAT")
    assert abs(pnl[1].value - base["cashflow"].pnl["pat"]) < 1e-9
    dcf = _find(wb["Cash Flow"], "DEVELOPER NET CASH FLOW (pre-tax)")
    total = dcf[-1].value if dcf[-1].value is not None else [c.value for c in dcf if c.value is not None][-1]
    assert abs(total - base["cashflow"].pnl["ebitda"]) < 1e-6           # cash = profit, also in the report
    irr = _find(wb["P&L & Metrics"], "Developer IRR (pre-tax)")
    assert irr[1].value == "110.2%" and irr[3].value == "Strong"
    assert wb["Summary"]["A4"].value == "✔ ALL CHECKS PASSED"


def test_report_can_be_reopened_as_the_same_project(project, base):
    assert to_dict(project_from_excel(build_excel_report(project, base))) == to_dict(project)


def test_report_with_stress_and_monte_carlo(project, base):
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    pytest.importorskip("plotly")
    import stress_view
    mc = monte_carlo(project, n=25, seed=3)
    wb = load_workbook(io.BytesIO(build_excel_report(project, base, stress_view.compute_tables(project), mc)))
    best = _find(wb["Stress Test"], "Best Case")
    assert abs(best[4].value - 1.5076) < 1e-3
    runs = [r for r in wb["Monte Carlo"].iter_rows(values_only=True) if isinstance(r[0], int) and len(r) >= 8 and r[7] is not None]
    assert len(runs) == 25


def test_report_works_for_other_deals():
    for change in (lambda p: setattr(p.info, "deal_type", "Outright"), lambda p: setattr(p.info, "project_type", "Villa"),
                   lambda p: setattr(p.info, "scenario", "Worst Case")):
        p = variant(change)
        assert len(build_excel_report(p, run_model(p))) > 10_000
